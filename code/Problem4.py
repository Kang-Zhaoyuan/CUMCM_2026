"""问题4：考虑径向收缩的温度—含水率耦合烘干。"""
from pathlib import Path
import argparse, json
from functools import partial
import numpy as np
from openpyxl import load_workbook
from scipy.integrate import solve_ivp
from scipy.interpolate import PchipInterpolator
from grid import Grid, R, Constant, DryEvent

ROOT=Path(__file__).resolve().parent.parent


def solve(nr=120,nz=120,rtol=1e-7,atol=1e-9,stretch=2.,end=259200,tail_points=20,method='BDF',output=None):
    data=[]
    for name in ('附件1.xlsx','附件2.xlsx'):
        book=load_workbook(ROOT/'problem'/'附件'/name,read_only=True,data_only=True)
        values=np.asarray(list(book.active.values)[1:],float); book.close()
        if not np.isfinite(values).all() or np.any(np.diff(values[:,0])<=0): raise ValueError(f'{name}数据无效')
        data.append(values)
    env,radius_data=data
    if not 0<end<=radius_data[-1,0]: raise ValueError('end须在附件2时间范围内')
    if np.any(radius_data[:,1]<=0): raise ValueError('附件2半径必须为正数')
    if not isinstance(tail_points,(int,np.integer)) or not 1<=tail_points<=len(env): raise ValueError(f'tail_points须为1到{len(env)}之间的整数')
    grid=Grid(nr,nz,stretch,shrinking=True)
    curve=PchipInterpolator(env[:,0],env[:,1:],axis=0); radius=PchipInterpolator(radius_data[:,0],radius_data[:,1]*.01)
    tail=tuple(np.mean(env[-tail_points:,1:],axis=0)); y=np.r_[np.full(grid.n,28.),np.full(grid.n,2.55)]
    times,radii=np.arange(0,end+1,60),np.linspace(0,R,21)
    tout,temp,moist=[0.],[np.full(21,28.)],[np.full(21,2.55)]
    surface_moist,current_radius=[2.55],[R*100]
    stats=dict(nr=nr,nz=nz,stretch=stretch,rtol=rtol,atol=atol,method=method,tail=list(tail),tail_points=tail_points,
               shrinking=True,steps=0,nfev=0,njev=0,nlu=0)
    breaks=np.unique(np.r_[0,env[(env[:,0]>0)&(env[:,0]<end),0],radius_data[(radius_data[:,0]>0)&(radius_data[:,0]<end),0],end])

    dry=DryEvent(grid.n)
    for a,b in zip(breaks[:-1],breaks[1:]):
        air=Constant(tail) if a>=env[-1,0] else curve
        fun=partial(grid.system,air=air,radius=radius)
        jac=partial(grid.system,air=air,jac=True,radius=radius)
        sol=solve_ivp(fun,(a,b),y,method=method,jac=jac,rtol=rtol,atol=atol,
                      dense_output=True,events=dry)
        if not sol.success: raise RuntimeError(sol.message)
        stop=sol.t[-1]; ts=times[(times>a)&(times<=stop)]
        if sol.status==1 or b==end: ts=np.unique(np.r_[ts,stop])
        for batch in np.array_split(ts,max(1,int(np.ceil(len(ts)/100)))):
            if not len(batch): continue
            for time,field in zip(batch,sol.sol(batch).T.reshape(-1,2,*grid.shape)):
                grid.geometry(float(radius(time))); tout.append(float(time))
                temp.append(PchipInterpolator(grid.r,field[0,:,0],extrapolate=False)(radii))
                moist.append(PchipInterpolator(grid.r,field[1,:,0],extrapolate=False)(radii))
                surface_moist.append(float(field[1,-1,0])); current_radius.append(float(grid.r[-1]*100))
        y=sol.y[:,-1]; stats['steps']+=len(sol.t)-1
        for key in ('nfev','njev','nlu'): stats[key]+=getattr(sol,key)
        if (b>=14400 and b%21600==0) or sol.status==1: print(f't={stop/3600:.4f} h, max C={y[grid.n:].max():.7f}',flush=True)
        if sol.status==1: break
    stats.update(dry_time_s=float(stop) if sol.status==1 else None,max_location=list(map(int,np.unravel_index(np.argmax(y[grid.n:]),grid.shape))))
    result=dict(time=np.array(tout),radius_cm=radii*100,temperature=np.array(temp),moisture=np.array(moist),
                surface_moisture=np.array(surface_moist),current_radius_cm=np.array(current_radius),
                final_temperature=y[:grid.n].reshape(grid.shape),final_moisture=y[grid.n:].reshape(grid.shape),r=grid.r,z=grid.z,xi=grid.r0/R)
    if output is not None:
        path=Path(output); book=load_workbook(ROOT/'problem'/'附件'/'附件3'/'result4.xlsx'); sheet=book.worksheets[0]
        if sheet.max_row>1: sheet.delete_rows(2,sheet.max_row-1)
        for col in range(2,sheet.max_column+1): sheet.cell(1,col).value=None
        headers=[*result['radius_cm'],'药材表面','半径/cm']
        for col,value in enumerate(headers,2): sheet.cell(1,col,round(float(value),1) if isinstance(value,(int,float,np.number)) else value)
        values=np.c_[result['moisture'][1:],result['surface_moisture'][1:],result['current_radius_cm'][1:]]
        for time,row in zip(result['time'][1:],values): sheet.append([float(time),*[round(float(v),4) if np.isfinite(v) else None for v in row]])
        for row in sheet.iter_rows(min_row=2,min_col=2):
            for cell in row: cell.number_format='0.0000'
        sheet.column_dimensions['A'].width=26; sheet.freeze_panes='B2'
        path.parent.mkdir(parents=True,exist_ok=True); book.save(path); book.close()
        print(f'第4问结果：{path}（{len(result["time"])-1}个时刻）')
    return result,stats


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    for name,kind,default in (('nr',int,120),('nz',int,120),('rtol',float,1e-7),('atol',float,1e-9),('stretch',float,2.)):
        parser.add_argument('--'+name,type=kind,default=default)
    parser.add_argument('--end',type=int,default=259200); parser.add_argument('--tail-points',type=int,default=20)
    parser.add_argument('--output',type=Path,default=ROOT/'problem'/'附件'/'附件3'/'result4.xlsx')
    _,stats=solve(**vars(parser.parse_args()))
    print('在指定时间上限内尚未达标。' if stats['dry_time_s'] is None else f'烘干时间：{stats["dry_time_s"]/3600:.4f} h')
    print(json.dumps(stats,indent=2))

