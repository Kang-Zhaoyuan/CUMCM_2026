"""问题2：温度—含水率耦合烘干。"""
from pathlib import Path
import argparse, json
from functools import partial
import numpy as np
from openpyxl import load_workbook
from scipy.integrate import solve_ivp
from scipy.interpolate import PchipInterpolator
from grid import Grid, R

ROOT=Path(__file__).resolve().parent.parent


def solve(nr=120,nz=120,rtol=1e-7,atol=1e-9,stretch=2.,end=10800,method='BDF',output=None):
    if not 0<end<=14400: raise ValueError('end须在附件1时间范围内')
    book=load_workbook(ROOT/'problem'/'附件'/'附件1.xlsx',read_only=True,data_only=True)
    env=np.asarray(list(book.active.values)[1:],float); book.close()
    if not np.isfinite(env).all() or np.any(np.diff(env[:,0])<=0): raise ValueError('附件1数据无效')
    grid,air=Grid(nr,nz,stretch),PchipInterpolator(env[:,0],env[:,1:],axis=0)
    y=np.r_[np.full(grid.n,28.),np.full(grid.n,2.55)]
    times,radii=np.arange(0,end+1),np.linspace(0,R,21)
    tout,temp,moist=[0.],[np.full(21,28.)],[np.full(21,2.55)]
    stats=dict(nr=nr,nz=nz,stretch=stretch,rtol=rtol,atol=atol,method=method,tail=list(env[-1,1:]),tail_points=None,
               shrinking=False,steps=0,nfev=0,njev=0,nlu=0)
    breaks=np.unique(np.r_[0,env[(env[:,0]>0)&(env[:,0]<end),0],end])
    fun,jac=partial(grid.system,air=air),partial(grid.system,air=air,jac=True)
    for a,b in zip(breaks[:-1],breaks[1:]):
        sol=solve_ivp(fun,(a,b),y,method=method,jac=jac,rtol=rtol,atol=atol,dense_output=True)
        if not sol.success: raise RuntimeError(sol.message)
        ts=times[(times>a)&(times<=b)]
        for time,field in zip(ts,sol.sol(ts).T.reshape(-1,2,*grid.shape)):
            tout.append(float(time)); temp.append(PchipInterpolator(grid.r,field[0,:,0])(radii)); moist.append(PchipInterpolator(grid.r,field[1,:,0])(radii))
        y=sol.y[:,-1]; stats['steps']+=len(sol.t)-1
        for key in ('nfev','njev','nlu'): stats[key]+=getattr(sol,key)
    stats.update(dry_time_s=None,max_location=list(map(int,np.unravel_index(np.argmax(y[grid.n:]),grid.shape))))
    result=dict(time=np.array(tout),radius_cm=radii*100,temperature=np.array(temp),moisture=np.array(moist),
                final_temperature=y[:grid.n].reshape(grid.shape),final_moisture=y[grid.n:].reshape(grid.shape),r=grid.r,z=grid.z)
    if output is not None:
        path=Path(output); book=load_workbook(ROOT/'problem'/'附件'/'附件3'/'result2.xlsx')
        for field,name in (('temperature','温度'),('moisture','水分浓度')):
            sheet=book[name]
            if sheet.max_row>1: sheet.delete_rows(2,sheet.max_row-1)
            for col in range(2,sheet.max_column+1): sheet.cell(1,col).value=None
            for col,value in enumerate(result['radius_cm'],2): sheet.cell(1,col,round(float(value),1))
            for time,row in zip(result['time'][1:],result[field][1:]): sheet.append([float(time),*[round(float(v),4) for v in row]])
            for row in sheet.iter_rows(min_row=2,min_col=2):
                for cell in row: cell.number_format='0.0000'
            sheet.column_dimensions['A'].width=26; sheet.freeze_panes='B2'
        path.parent.mkdir(parents=True,exist_ok=True); book.save(path); book.close()
        print(f'第2问结果：{path}（{len(result["time"])-1}个时刻）')
    return result,stats


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    for name,kind,default in (('nr',int,120),('nz',int,120),('rtol',float,1e-7),('atol',float,1e-9),('stretch',float,2.)):
        parser.add_argument('--'+name,type=kind,default=default)
    parser.add_argument('--end',type=int,default=10800)
    parser.add_argument('--output',type=Path,default=ROOT/'problem'/'附件'/'附件3'/'result2.xlsx')
    _,stats=solve(**vars(parser.parse_args())); print(json.dumps(stats,indent=2))

